"""Build the SiNK 2026 talk from the validated repo figures.

Requires python-pptx and Pillow. Run after generate_conference_figures.py.
The first 12 slides are the oral talk; eight additional slides hold remaining figures.
"""
from __future__ import annotations

from pathlib import Path
import json

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "research" / "conference_figures"
LOGO = ROOT / "research" / "assets" / "rvce_logo.png"
OUTPUT = ROOT / "SINk2026_Phase1_CV_Digital_Twin.pptx"

NAVY = RGBColor(22, 48, 75)
BLUE = RGBColor(26, 110, 167)
TEAL = RGBColor(17, 133, 128)
CORAL = RGBColor(204, 83, 65)
GRAY = RGBColor(78, 97, 113)
PALE = RGBColor(225, 235, 241)
WHITE = RGBColor(255, 255, 255)
FONT = "Aptos"


def add_text(slide, text, x, y, w, h, size=20, color=NAVY, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return shape


def add_paragraphs(slide, lines, x, y, w, h, size=20, spacing=16, color=NAVY):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
    return shape


def add_logo(slide):
    # Retain the original artwork's aspect ratio and white ground.
    slide.shapes.add_picture(str(LOGO), Inches(10.73), Inches(.20), width=Inches(2.2))


def base_slide(ppt, title, number, *, section="RESEARCH PRESENTATION"):
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_logo(slide)
    if title:
        add_text(slide, title, .68, .43, 9.88, .73, 31, NAVY, True)
    add_text(slide, f"SiNK 2026  ·  {section}", .68, 7.13, 9.3, .22, 10, GRAY)
    add_text(slide, str(number).zfill(2), 12.22, 7.10, .43, .25, 10, GRAY, align=PP_ALIGN.RIGHT)
    return slide


def add_contained(slide, path, x, y, w, h):
    with Image.open(path) as im:
        aspect = im.width / im.height
    if aspect > w / h:
        image_h = w / aspect
        x1, y1, w1, h1 = x, y + (h - image_h) / 2, w, image_h
    else:
        image_w = h * aspect
        x1, y1, w1, h1 = x + (w - image_w) / 2, y, image_w, h
    slide.shapes.add_picture(str(path), Inches(x1), Inches(y1), width=Inches(w1), height=Inches(h1))


def note(slide, body):
    slide.notes_slide.notes_text_frame.text = body + "\nLogo source: https://d2lk14jtvqry1q.cloudfront.net/media/large_89_9b01402b5b_926235f86d.png"


def main():
    for needed in [LOGO, *(FIG / f for f in [
        "00_digital_twin_roadmap.png", "00_measurement_workflow.png", "01_recorded_cv.png",
        "04_boundary_extrema.png", "08_simulated_noise.png", "11_synthetic_split_comparison.png",
    ])]:
        if not needed.exists():
            raise FileNotFoundError(needed)

    ppt = Presentation()
    ppt.slide_width = Inches(13.333)
    ppt.slide_height = Inches(7.5)

    s = base_slide(ppt, "", 1)
    add_text(s, "Toward a battery digital twin", .80, 1.55, 11.6, .74, 43, NAVY, True)
    add_text(s, "Phase 1: reliable cyclic voltammetry evidence", .82, 2.38, 11.6, .52, 28, BLUE)
    add_text(s, "Cheritha Kaiwar  ·  Prajwal S  ·  Karan Koder", .83, 4.35, 11.4, .42, 23, NAVY)
    add_text(s, "Under the guidance of Dr. Manjunatha C.", .83, 4.90, 11.4, .39, 20, GRAY)
    add_text(s, "RV College of Engineering, Bengaluru", .83, 5.62, 11.4, .35, 17, BLUE)
    note(s, "Title and authors supplied by the project team. Scope: implemented CV quality layer, with future battery-health validation planned.")

    s = base_slide(ppt, "Problem statement", 2)
    add_text(s, "Industry needs battery diagnostics grounded in reliable measurements.", .82, 1.53, 11.4, 1.1, 30, NAVY, True)
    add_paragraphs(s, [
        "A CV trace can have missing units or acquisition details.",
        "A scan boundary can cut off a peak and make its potential unknowable.",
        "A health model trained on generated labels cannot establish physical-cell accuracy.",
    ], .86, 3.02, 11.5, 2.5, 21, 20)
    add_text(s, "Engineering need: check the evidence before reporting features or health estimates.", .84, 6.20, 11.7, .52, 21, TEAL, True)
    note(s, "Industrial context: NREL Battery Lifespan research, https://www.nrel.gov/transportation/battery-lifespan.html. CV and model limitations are findings from this repository.")

    s = base_slide(ppt, "Existing solutions and our contribution", 3)
    rows = [
        ("Battery management", "Uses operating voltage, current and temperature to estimate state."),
        ("Potentiostat software", "Plots CV scans and provides peak and baseline analysis."),
        ("Health prediction research", "Trains models on features paired with measured capacity."),
        ("This Phase 1 prototype", "Keeps CV acquisition order, checks support for features, and shows why a value is withheld."),
    ]
    tbl = s.shapes.add_table(4, 2, Inches(.77), Inches(1.55), Inches(11.74), Inches(4.95)).table
    tbl.columns[0].width = Inches(3.47)
    tbl.columns[1].width = Inches(8.27)
    for i, (name, description) in enumerate(rows):
        for j, value in enumerate((name, description)):
            c = tbl.cell(i, j)
            c.text = value
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(242, 247, 249)
            c.margin_left = c.margin_right = Inches(.15)
            c.margin_top = c.margin_bottom = Inches(.09)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in c.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(20 if j == 0 else 18)
                p.font.bold = j == 0
                p.font.color.rgb = NAVY if j == 0 else GRAY
    note(s, "Sources: NREL https://www.nrel.gov/transportation/battery-lifespan.html ; Gamry Echem Analyst manual https://www.gamry.com/assets/Uploads/EchemAnalystSoftwareManual.pdf ; battery SoH ML on labeled cells https://arxiv.org/abs/2102.00837 . No claim of unique invention or superiority over existing CV tools.")

    s = base_slide(ppt, "Research question and scope", 4)
    add_text(s, "Detect when a CV scan does not support a reported feature", .80, 1.44, 11.7, .88, 29, NAVY, True)
    add_contained(s, FIG / "00_digital_twin_roadmap.png", 1.25, 2.35, 10.84, 4.43)
    note(s, "Phase 1 is implemented in the repo. Phase 2 requires repeated measurements and independent capacity labels. Phase 3 requires a calibrated model that updates for an identified physical battery. Figure is generated from research/generate_conference_figures.py. SINk 2026 organiser theme context: https://www.linkedin.com/posts/dr-vishal-chaudhary_schedule-activity-7508161441263304704-ga_Y . The fit is electrochemical sensing linked to interpretable decisions, with an energy application.")

    s = base_slide(ppt, "Methodology: CV quality gate", 5)
    add_contained(s, FIG / "00_measurement_workflow.png", .74, 1.39, 11.85, 5.67)
    note(s, "Data flow: preserve ordered CSV; require explicit units in the live API; check forward and return branch, closure and interior extrema; report only supported descriptors. The synthetic RF model is a separate dataset and not connected to the measured trace. Figure generated from repository source.")

    s = base_slide(ppt, "Experimental input: one reported CV sweep", 6)
    add_contained(s, FIG / "01_recorded_cv.png", .57, 1.55, 8.13, 5.40)
    add_text(s, "21 measured values", 8.96, 1.86, 3.52, .42, 23, BLUE, True)
    add_paragraphs(s, [
        "Team reports a real CV setup.",
        "Approximately room temperature.",
        "Electrodes reported: graphite, platinum and Ag/AgCl. Roles not documented.",
        "Scan rate and current unit not recorded in the CSV.",
    ], 8.97, 2.52, 3.48, 3.30, 18, 11)
    note(s, "Source: attached_assets/14,12csv_1765724360390.csv and research/experimental_metadata.json. Current is plotted in µA only under a provisional assumption that the source numbers are amperes. Team reports experimental origin, but original instrument export is unavailable.")

    s = base_slide(ppt, "Measured result: peak separation withheld", 7)
    add_contained(s, FIG / "04_boundary_extrema.png", .68, 1.48, 9.06, 5.45)
    add_text(s, "Quality: limited", 9.86, 2.18, 2.66, .45, 23, CORAL, True)
    add_paragraphs(s, [
        "Apparent extrema sit at the scan limits.",
        "The true peak positions may lie beyond the voltage window.",
        "The system leaves ΔEp blank and gives the reason.",
    ], 9.86, 2.90, 2.62, 2.71, 18, 10)
    note(s, "Source: supplied 21-point CSV, analyzed by research/cv_sensor_study.py and the live API. The signed path integral is approximately -1.75e-7 A·V under the provisional ampere assumption; it is not charge, energy or capacity.")

    s = base_slide(ppt, "Controlled test: sensitivity to current noise", 8)
    add_contained(s, FIG / "08_simulated_noise.png", .65, 1.47, 9.10, 5.46)
    add_text(s, "At 10% noise", 9.91, 2.08, 2.50, .40, 22, BLUE, True)
    add_text(s, "10 mV", 9.91, 2.70, 2.50, .52, 29, NAVY, True)
    add_text(s, "median absolute error", 9.91, 3.20, 2.50, .32, 16, GRAY)
    add_text(s, "30.5 mV", 9.91, 4.02, 2.50, .48, 27, NAVY, True)
    add_text(s, "95th percentile error", 9.91, 4.52, 2.50, .32, 16, GRAY)
    add_text(s, "Simulation only; 200 replicates per level", 9.91, 5.44, 2.50, .64, 16, CORAL)
    note(s, "Source: research/results/sensor_report.json, seed 20260926. Gaussian peak simulation: +/-0.10 V peaks, 3 microamp amplitude, 10 mV sampling. Independent Gaussian current noise. At 10% current noise: median absolute separation error 10 mV, empirical 95th percentile 30.5 mV across 200 seeded replicates. This measures software extractor sensitivity under one signal model. It is not an instrument error specification or physical battery accuracy. The single experimental scan has 50 mV potential increments and no reference peak or repeats, so physical sensing accuracy is unknown.")

    metrics = json.loads((FIG / "figure_metrics.json").read_text(encoding="utf-8"))["model_split_audit"]
    random_metrics = metrics["Random row split"]
    late_metrics = metrics["Later cycles held out"]
    s = base_slide(ppt, "Random Forest regression metrics", 9)
    add_text(s, f"Generated-data random split: R² {random_metrics['r2']:.4f}, RMSE {random_metrics['rmse_soh_points']:.2f} points, MAE {random_metrics['mae_soh_points']:.2f} points", .78, 1.35, 11.68, .43, 20, BLUE, True)
    add_contained(s, FIG / "11_synthetic_split_comparison.png", .70, 1.83, 11.85, 5.10)
    note(s, f"Source: ml/training_data.csv and research/conference_figures/figure_metrics.json. Features were generated from the synthetic SoH trajectory. Random row split: RMSE {random_metrics['rmse_soh_points']:.4f} SoH points, MAE {random_metrics['mae_soh_points']:.4f} points, R2 {random_metrics['r2']:.6f}. Last 30 cycles held out: RMSE {late_metrics['rmse_soh_points']:.4f} points, MAE {late_metrics['mae_soh_points']:.4f} points, R2 {late_metrics['r2']:.6f}. R2 is a regression fit measure, not classification accuracy. Neither test validates a physical cell. The live API does not run this model.")

    s = base_slide(ppt, "Conclusion", 10)
    add_text(s, "Phase 1 provides a working CV evidence layer.", .84, 1.52, 11.65, .74, 31, NAVY, True)
    add_paragraphs(s, [
        "The prototype preserves the reported experimental trace and checks sweep quality.",
        "It correctly withholds peak separation for boundary-limited data.",
        "The simulation tests extraction sensitivity; the synthetic model audit does not validate SoH.",
    ], .89, 2.74, 11.60, 3.07, 22, 21)
    note(s, "Conclusion limited to implemented software behavior and one team-reported experimental trace. The project has no independently measured capacity labels, so it cannot claim health prediction accuracy or a complete digital twin.")

    s = base_slide(ppt, "Future scope", 11)
    add_text(s, "Phase 2: physical validation", .83, 1.51, 11.65, .44, 25, BLUE, True)
    add_paragraphs(s, [
        "Recover instrument exports, electrode roles, material identity, units and scan rate.",
        "Repeat CV on identified cells and pair each sweep with measured discharge capacity.",
        "Test repeatability and any SoH model on entire cells held out from training.",
    ], .87, 2.11, 11.45, 2.47, 20, 16)
    add_text(s, "Phase 3: operational twin", .83, 5.02, 11.65, .44, 25, TEAL, True)
    add_text(s, "Calibrate a battery-specific state model, update it from new data, and verify it over time.", .87, 5.64, 11.42, .83, 20, NAVY)
    note(s, "Future work proposal, not achieved results. The original 21-point CSV has no capacity target or complete acquisition metadata. Proposed physical validation follows the principles of separated cells and independently measured outcomes.")

    s = base_slide(ppt, "References", 12)
    refs = [
        "Project data and methods. Battery-digital-twin repository: experimental CSV, metadata record, sensor study, and synthetic audit (2026).",
        "NREL. Battery Lifespan: diagnostics and model-based battery health research. nrel.gov/transportation/battery-lifespan.html",
        "Elgrishi et al. A Practical Beginner’s Guide to Cyclic Voltammetry. Journal of Chemical Education 95 (2018). DOI: 10.1021/acs.jchemed.7b00361",
        "Gamry Instruments. Echem Analyst software manual: cyclic voltammetry analysis. gamry.com/assets/Uploads/EchemAnalystSoftwareManual.pdf",
        """"Battery digital twins: Perspectives on the fusion of models, data and artificial intelligence." Energy and AI 1 (2020). DOI: 10.1016/j.egyai.2020.100016""",
        """"Machine learning pipeline for battery state of health estimation" (2021). arXiv:2102.00837""",
    ]
    add_paragraphs(s, refs, .81, 1.44, 11.65, 5.63, 16, 13)
    note(s, "Full URLs: https://www.nrel.gov/transportation/battery-lifespan.html ; https://pubs.acs.org/doi/10.1021/acs.jchemed.7b00361 ; https://www.gamry.com/assets/Uploads/EchemAnalystSoftwareManual.pdf ; https://doi.org/10.1016/j.egyai.2020.100016 ; https://arxiv.org/abs/2102.00837 .")

    appendix = [
        ("Electrode materials reported by the team", "00_electrode_materials.png", "The team reported graphite, platinum and Ag/AgCl. Exact roles, material and electrolyte are not documented."),
        ("Potential sequence", "02_potential_sequence.png", "Same team-reported 21-point experimental trace. Time and scan rate unavailable."),
        ("Current sequence", "03_current_sequence.png", "Same team-reported experimental trace; current unit provisional."),
        ("Forward and return current difference", "05_branch_difference.png", "Derived from the same 21 points. Descriptive curve, no mechanism attribution."),
        ("Potential sampling steps", "06_voltage_steps.png", "Step size cannot determine scan rate without timestamps."),
        ("Signed path integral", "07_cumulative_path_integral.png", "A·V under the provisional ampere interpretation. This is not charge or energy."),
        ("Generated SoH trajectory", "09_synthetic_soh_trajectory.png", "From ml/training_data.csv; no measured capacity values."),
        ("Synthetic feature correlations", "10_synthetic_feature_correlations.png", "Features derive from generated SoH labels; correlations are not physical discovery."),
    ]
    for idx, (title, filename, detail) in enumerate(appendix, 13):
        s = base_slide(ppt, title, idx, section="FIGURE APPENDIX")
        add_contained(s, FIG / filename, .70, 1.45, 11.87, 5.55)
        note(s, detail + " Figure generated by research/generate_conference_figures.py from repository data.")

    ppt.save(OUTPUT)
    print(f"Created {OUTPUT} with {len(ppt.slides)} slides (12 talk + 8 appendix)")


if __name__ == "__main__":
    main()
